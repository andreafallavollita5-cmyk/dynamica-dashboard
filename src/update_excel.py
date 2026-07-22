"""Generate the client-facing Excel snapshot from the final report DataFrame."""

from __future__ import annotations

import calendar
import json
import math
import os
import shutil
import tempfile
import textwrap
from copy import copy
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.report_groups import CLIENT_GROUPS, SUMMARY_FIELDS, client_subtotal_group
from src.combined_campaigns import COMBINED_CAMPAIGN_PAIRS
from src.dates import weekdays_in_month


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_TEMPLATE = ROOT / "templates/Template_export_cliente_Dynamica.xlsx"
DEFAULT_OUTPUT_DIR = ROOT / "exports"
LATEST_FILENAME = "report_dynamica_updated.xlsx"

REPORT_SHEET = "Report Cliente"
DATA_SHEET = "Dati"
CONFIG_SHEET = "Config"
FIRST_DATA_ROW = 5
TEMPLATE_CAPACITY = 30

REPORT_FIELDS = (
    "funnel",
    "platform",
    "channel",
    "campaign_id",
    "campaign_name",
    "investimento_media",
    "percentuale_investimento",
    "cpp_medio",
    "stima_pratiche",
    "cpl_target",
    "stima_lead",
    "stima_lead_giornaliere",
    "stima_lead_progressiva",
    "lead_effettive",
    "delta_lead",
    "stima_spending_giornaliera",
    "stima_spending_progressiva",
    "speso_effettivo",
    "delta_speso",
    "delta_delivery_pct",
    "cpl_effettivo",
    "delta_cpl",
    "action",
)

CURRENCY_COLUMNS = (6, 8, 10, 16, 17, 18, 19, 21, 22)
PERCENT_COLUMNS = (7, 20)
NUMBER_COLUMNS = (9, 11, 12, 13, 14, 15)
CURRENCY_FORMAT = '#,##0.00 "€"'
PERCENT_FORMAT = "0.0%"
NUMBER_FORMAT = "0.0"
DATE_FORMAT = "dd/mm/yyyy"
DATETIME_FORMAT = "dd/mm/yyyy hh:mm"

LIGHT_BLUE = "DCE6F1"
BLACK = "000000"
WHITE = "FFFFFF"
GREEN = "C6EFCE"
GREEN_TEXT = "006100"
RED = "FFC7CE"
RED_TEXT = "9C0006"
YELLOW = "FFEB9C"
YELLOW_TEXT = "9C6500"


def _resolve_path(path: str | os.PathLike[str], default: Path) -> Path:
    candidate = Path(path)
    if candidate == Path(str(default.relative_to(ROOT))):
        return default
    return candidate if candidate.is_absolute() else ROOT / candidate


def _is_blank(value: Any) -> bool:
    if value is None:
        return True
    try:
        return bool(pd.isna(value))
    except (TypeError, ValueError):
        return False


def _as_number(value: Any) -> float | None:
    if _is_blank(value) or str(value).strip() == "":
        return None
    try:
        numeric = float(value)
    except (TypeError, ValueError):
        return None
    return numeric if math.isfinite(numeric) else None


def _as_text(value: Any) -> str:
    return "" if _is_blank(value) else str(value).strip()


def _as_action_text(value: Any) -> str:
    """Preserve the manual Action content exactly as supplied."""
    return "" if _is_blank(value) else str(value)


def _as_datetime(value: Any) -> datetime | None:
    if _is_blank(value) or str(value).strip() == "":
        return None
    parsed = pd.to_datetime(value, errors="coerce")
    if pd.isna(parsed):
        return None
    if getattr(parsed, "tzinfo", None) is not None:
        parsed = parsed.tz_localize(None)
    return parsed.to_pydatetime()


def _safe_divide(numerator: Any, denominator: Any) -> float | None:
    numerator_value = _as_number(numerator)
    denominator_value = _as_number(denominator)
    if numerator_value is None or denominator_value in (None, 0):
        return None
    return numerator_value / denominator_value


def _wrapped_line_count(value: Any, width: int = 34) -> int:
    """Estimate Excel wrapped lines conservatively for fixed-width text columns."""
    text = _as_action_text(value)
    if not text:
        return 1
    lines = 0
    for paragraph in text.splitlines() or [text]:
        wrapped = textwrap.wrap(
            paragraph,
            width=width,
            break_long_words=False,
            break_on_hyphens=False,
            replace_whitespace=False,
        )
        lines += max(1, len(wrapped))
    return lines


def _enabled(value: Any) -> bool:
    if _is_blank(value):
        return False
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value != 0
    return str(value).strip().casefold() in {"true", "1", "yes", "y", "si", "sì"}


def _classify_source_status(report_df: pd.DataFrame) -> str:
    statuses = " ".join(
        _as_text(value).casefold()
        for value in report_df.get("source_status", pd.Series(dtype="object"))
        if not _is_blank(value)
    )
    if "mock" in statuses or "demo" in statuses:
        return "mock"

    failure_markers = (
        ":error",
        "failed",
        "failure",
        "unavailable",
        "ambiguous",
        "unauthorized",
    )
    has_failure = any(marker in statuses for marker in failure_markers)
    useful_columns = [
        column
        for column in (
            "campaign_name",
            "investimento_media",
            "stima_lead",
            "lead_effettive",
            "speso_effettivo",
        )
        if column in report_df.columns
    ]
    has_usable_data = not report_df.empty and any(
        report_df[column].map(lambda value: not _is_blank(value)).any()
        for column in useful_columns
    )
    if has_failure:
        return "partial" if has_usable_data else "error"
    return "success" if has_usable_data else "error"


def _prepare_dataframe(report_df: pd.DataFrame) -> pd.DataFrame:
    frame = report_df.copy()
    if "row_type" not in frame.columns:
        frame["row_type"] = "campaign"
    else:
        frame["row_type"] = frame["row_type"].fillna("campaign")
    if "enabled" in frame.columns:
        frame = frame[frame["enabled"].map(_enabled)].copy()

    frame["_funnel_order"] = pd.factorize(frame.get("funnel", ""), sort=False)[0]
    frame["_original_order"] = range(len(frame))
    frame["_platform_sort"] = frame.get("platform", "").fillna("").astype(str).str.casefold()
    frame["_excel_row_sort"] = pd.to_numeric(
        frame.get("excel_row"), errors="coerce"
    ).fillna(float("inf"))
    frame["_campaign_sort"] = frame.get("campaign_name", "").fillna("").astype(str).str.casefold()
    frame = frame.sort_values(
        [
            "_funnel_order",
            "_platform_sort",
            "_excel_row_sort",
            "_campaign_sort",
            "_original_order",
        ],
        kind="stable",
    )
    return frame.drop(
        columns=[
            "_funnel_order",
            "_original_order",
            "_platform_sort",
            "_excel_row_sort",
            "_campaign_sort",
        ]
    ).reset_index(drop=True)


def _copy_row_style(ws, source_row: int, target_row: int) -> None:
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height
    for column in range(1, ws.max_column + 1):
        source = ws.cell(source_row, column)
        target = ws.cell(target_row, column)
        target._style = copy(source._style)
        target.alignment = copy(source.alignment)
        target.number_format = source.number_format
        target.protection = copy(source.protection)


def _merge_combined_campaign_metrics(ws, output_rows: dict[int, int]) -> None:
    """Merge planned/delta/CPL cells while leaving effective leads separate."""
    for parent_excel_row, child_excel_row in COMBINED_CAMPAIGN_PAIRS:
        parent = output_rows.get(parent_excel_row)
        child = output_rows.get(child_excel_row)
        if parent is None or child is None:
            continue
        if child != parent + 1:
            raise ValueError(
                "Le campagne CRM combinate non sono adiacenti nell'export Excel."
            )
        for column in (11, 12, 13, 15, 21, 22):
            ws.merge_cells(
                start_row=parent,
                start_column=column,
                end_row=child,
                end_column=column,
            )
            ws.cell(parent, column).alignment = Alignment(
                horizontal=ws.cell(parent, column).alignment.horizontal,
                vertical="center",
                wrap_text=ws.cell(parent, column).alignment.wrap_text,
            )


def _campaign_id(row: pd.Series) -> str:
    for field in ("google_campaign_id", "meta_campaign_id", "dynamics_campaign_key"):
        value = _as_text(row.get(field))
        if value:
            return value
    return ""


def _client_subtotal_groups(frame: pd.DataFrame) -> list[tuple[str, str, pd.DataFrame]]:
    labels = frame.apply(lambda row: client_subtotal_group(row.to_dict()), axis=1)
    return [
        (slug, label, frame[labels == slug])
        for slug, label in CLIENT_GROUPS
        if (labels == slug).any()
    ]


def _sum_values(frame: pd.DataFrame, column: str) -> float | None:
    if column not in frame.columns:
        return None
    values = [_as_number(value) for value in frame[column]]
    valid = [value for value in values if value is not None]
    return sum(valid) if valid else None


def _first_official_value(frame: pd.DataFrame, column: str) -> float | None:
    if column not in frame.columns:
        return None
    for value in frame[column]:
        number = _as_number(value)
        if number is not None:
            return number
    return None


def _official_summary_values(
    row: pd.Series,
    days_in_month: int,
    spending_daily_fallback: float | None = None,
) -> list[Any]:
    """Read one official report row; never rebuild subtotal or total values."""
    values: list[Any] = []
    for field in SUMMARY_FIELDS:
        if field == "stima_spending_giornaliera":
            value = _as_number(row.get(field))
            if value is None:
                value = (
                    spending_daily_fallback
                    if spending_daily_fallback is not None
                    else _safe_divide(row.get("investimento_media"), days_in_month)
                )
        else:
            value = _as_number(row.get(field))
        values.append(value)
    return values + [None]


def _legacy_summary_values(frame: pd.DataFrame, slug: str | None = None) -> list[Any]:
    """Compatibility path for historical in-memory frames used by older callers."""
    columns = (
        [f"kpi_{field}_totale" for field in SUMMARY_FIELDS]
        if slug is None
        else [f"subtotal_{slug}_{field}" for field in SUMMARY_FIELDS]
    )
    return [_first_official_value(frame, column) for column in columns] + [None]


def _write_report_values(ws, row_number: int, values: Iterable[Any]) -> None:
    for column, value in enumerate(values, start=1):
        ws.cell(row_number, column).value = value
    ws.cell(row_number, 4).number_format = "@"
    for column in CURRENCY_COLUMNS:
        ws.cell(row_number, column).number_format = CURRENCY_FORMAT
    for column in PERCENT_COLUMNS:
        ws.cell(row_number, column).number_format = PERCENT_FORMAT
    for column in NUMBER_COLUMNS:
        ws.cell(row_number, column).number_format = NUMBER_FORMAT
    ws.cell(row_number, 23).alignment = Alignment(wrap_text=True, vertical="top")


def _style_subtotal(ws, row_number: int) -> None:
    fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    top = Side(style="medium", color="7F8C8D")
    for cell in ws[row_number][:23]:
        cell.fill = fill
        cell.font = copy(cell.font)
        cell.font = Font(
            name=cell.font.name,
            size=cell.font.sz,
            bold=True,
            italic=cell.font.italic,
            color=cell.font.color,
        )
        cell.border = Border(
            left=cell.border.left,
            right=cell.border.right,
            top=top,
            bottom=cell.border.bottom,
        )


def _style_total(ws, row_number: int) -> None:
    for cell in ws[row_number][:23]:
        cell.fill = PatternFill("solid", fgColor=BLACK)
        cell.font = Font(
            name=cell.font.name or "Arial",
            size=cell.font.sz or 10,
            bold=True,
            color=WHITE,
        )
        cell.alignment = copy(cell.alignment)


def _add_conditional_formatting(ws, last_row: int) -> None:
    ws.conditional_formatting._cf_rules.clear()
    green_fill = PatternFill("solid", fgColor=GREEN)
    green_font = Font(color=GREEN_TEXT)
    red_fill = PatternFill("solid", fgColor=RED)
    red_font = Font(color=RED_TEXT)
    yellow_fill = PatternFill("solid", fgColor=YELLOW)
    yellow_font = Font(color=YELLOW_TEXT)

    # With delta_lead = effective - planned, positive values are favourable.
    ws.conditional_formatting.add(
        f"O{FIRST_DATA_ROW}:O{last_row}",
        CellIsRule(operator="greaterThan", formula=["0"], fill=green_fill, font=green_font),
    )
    ws.conditional_formatting.add(
        f"O{FIRST_DATA_ROW}:O{last_row}",
        CellIsRule(operator="lessThan", formula=["0"], fill=red_fill, font=red_font),
    )
    ws.conditional_formatting.add(
        f"T{FIRST_DATA_ROW}:T{last_row}",
        CellIsRule(operator="lessThan", formula=["-0.2"], fill=red_fill, font=red_font),
    )
    ws.conditional_formatting.add(
        f"T{FIRST_DATA_ROW}:T{last_row}",
        CellIsRule(
            operator="between",
            formula=["-0.2", "0"],
            fill=yellow_fill,
            font=yellow_font,
        ),
    )
    ws.conditional_formatting.add(
        f"T{FIRST_DATA_ROW}:T{last_row}",
        CellIsRule(operator="greaterThan", formula=["0"], fill=green_fill, font=green_font),
    )
    ws.conditional_formatting.add(
        f"V{FIRST_DATA_ROW}:V{last_row}",
        CellIsRule(operator="lessThan", formula=["0"], fill=green_fill, font=green_font),
    )
    ws.conditional_formatting.add(
        f"V{FIRST_DATA_ROW}:V{last_row}",
        CellIsRule(operator="greaterThan", formula=["0"], fill=red_fill, font=red_font),
    )


def _write_data_sheet(
    ws, frame: pd.DataFrame, days_in_month: int, dem_days_in_month: int
) -> None:
    for merged_range in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merged_range))
    ws.delete_rows(1, ws.max_row)

    data_frame = frame.copy()
    if "stima_spending_giornaliera" not in data_frame.columns:
        insert_at = (
            data_frame.columns.get_loc("stima_spending_progressiva")
            if "stima_spending_progressiva" in data_frame.columns
            else len(data_frame.columns)
        )
        daily = data_frame.apply(
            lambda row: _safe_divide(
                row.get("investimento_media"),
                dem_days_in_month
                if client_subtotal_group(row.to_dict()) == "dem"
                else days_in_month,
            ),
            axis=1,
        )
        data_frame.insert(insert_at, "stima_spending_giornaliera", daily)

    columns = list(data_frame.columns)
    for column_number, column_name in enumerate(columns, start=1):
        cell = ws.cell(1, column_number, column_name)
        cell.fill = PatternFill("solid", fgColor=BLACK)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    id_columns = {"google_campaign_id", "meta_campaign_id", "dynamics_campaign_key"}
    date_columns = {"report_date", "start_date", "end_date"}
    currency_fields = {
        "investimento_media",
        "cpp_medio",
        "cpl_target",
        "stima_spending_giornaliera",
        "stima_spending_progressiva",
        "speso_effettivo",
        "delta_speso",
        "cpl_effettivo",
        "delta_cpl",
    }
    percent_fields = {"percentuale_investimento", "delta_delivery_pct"}

    for row_number, (_, row) in enumerate(data_frame.iterrows(), start=2):
        for column_number, column_name in enumerate(columns, start=1):
            value = row.get(column_name)
            cell = ws.cell(row_number, column_number)
            if column_name in id_columns:
                cell.value = _as_text(value)
                cell.number_format = "@"
            elif column_name in date_columns:
                parsed = _as_datetime(value)
                cell.value = parsed.date() if parsed else None
                cell.number_format = DATE_FORMAT
            else:
                cell.value = None if _is_blank(value) else value
                if column_name in currency_fields:
                    cell.number_format = CURRENCY_FORMAT
                elif column_name in percent_fields:
                    cell.number_format = PERCENT_FORMAT

    last_row = max(1, len(data_frame) + 1)
    last_column = max(1, len(columns))
    ws.auto_filter.ref = f"A1:{get_column_letter(last_column)}{last_row}"
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    for column_number, column_name in enumerate(columns, start=1):
        width = 18
        if column_name in {"campaign_name", "source_status", "action"}:
            width = 35
        elif column_name.endswith("_id") or column_name.endswith("_key"):
            width = 23
        ws.column_dimensions[get_column_letter(column_number)].width = width


def _metadata(report_df: pd.DataFrame, frame: pd.DataFrame) -> dict[str, Any]:
    supplied = dict(report_df.attrs.get("metadata", {}))

    def first_value(column: str) -> Any:
        if column not in frame.columns:
            return None
        values = frame[column].dropna()
        return None if values.empty else values.iloc[0]

    start_date = _as_datetime(supplied.get("start_date") or first_value("start_date"))
    end_date = _as_datetime(supplied.get("end_date") or first_value("end_date"))
    updated_at = _as_datetime(supplied.get("updated_at") or first_value("report_date"))
    if start_date is None or end_date is None:
        raise ValueError("Il periodo del report non è disponibile.")
    if updated_at is None:
        updated_at = datetime.now()
    days_in_month = calendar.monthrange(start_date.year, start_date.month)[1]
    elapsed_days = max((end_date.date() - start_date.date()).days + 1, 0)
    return {
        "client_name": _as_text(supplied.get("client") or "Dynamica Retail"),
        "start_date": start_date,
        "end_date": end_date,
        "days_in_month": days_in_month,
        "elapsed_days": elapsed_days,
        "updated_at": updated_at,
        "source_status": _classify_source_status(frame),
    }


def _validate_workbook(path: Path, campaign_count: int, expected_end_date: date) -> None:
    workbook = load_workbook(path, data_only=False, read_only=False)
    try:
        if workbook.sheetnames != [REPORT_SHEET, DATA_SHEET, CONFIG_SHEET]:
            raise ValueError("Struttura fogli non valida.")
        if workbook[REPORT_SHEET].sheet_state != "visible":
            raise ValueError("Il report cliente non è visibile.")
        if workbook[DATA_SHEET].sheet_state != "hidden" or workbook[CONFIG_SHEET].sheet_state != "hidden":
            raise ValueError("I fogli tecnici non sono nascosti.")
        report_ws = workbook[REPORT_SHEET]
        total_rows = [
            row for row in range(FIRST_DATA_ROW, report_ws.max_row + 1)
            if report_ws.cell(row, 1).value == "TOTALE GENERALE"
        ]
        if len(total_rows) != 1:
            raise ValueError("Totale generale mancante o duplicato.")
        exported_campaigns = sum(
            1
            for row in range(FIRST_DATA_ROW, total_rows[0])
            if report_ws.cell(row, 5).value not in (None, "")
        )
        if exported_campaigns != campaign_count:
            raise ValueError("Numero campagne esportate non coerente.")
        config_end = workbook[CONFIG_SHEET]["B4"].value
        if getattr(config_end, "date", lambda: config_end)() != expected_end_date:
            raise ValueError("Data fine periodo non coerente.")
    finally:
        workbook.close()


def generate_client_excel(
    report_df: pd.DataFrame,
    template_path: str = "templates/Template_export_cliente_Dynamica.xlsx",
    output_dir: str = "exports",
) -> str:
    """Genera il report Excel cliente e restituisce il path del file creato."""
    template = _resolve_path(template_path, DEFAULT_TEMPLATE)
    destination_dir = _resolve_path(output_dir, DEFAULT_OUTPUT_DIR)
    if not template.exists():
        raise FileNotFoundError("Template Excel non disponibile.")

    frame = _prepare_dataframe(report_df)
    campaign_frame = frame[
        frame.get("row_type", "campaign").fillna("campaign") == "campaign"
    ].copy()
    official_frame = frame[frame.get("row_type", "") != "campaign"].copy()
    if campaign_frame.empty:
        raise ValueError("Nessuna campagna disponibile per l'export.")
    metadata = _metadata(report_df, campaign_frame)
    days_in_month = metadata["days_in_month"]
    dem_days_in_month = weekdays_in_month(metadata["start_date"].date())
    funnel_groups = _client_subtotal_groups(campaign_frame)
    campaign_daily_values = campaign_frame.apply(
        lambda row: (
            _as_number(row.get("stima_spending_giornaliera"))
            if _as_number(row.get("stima_spending_giornaliera")) is not None
            else _safe_divide(
                row.get("investimento_media"),
                dem_days_in_month
                if client_subtotal_group(row.to_dict()) == "dem"
                else days_in_month,
            )
        ),
        axis=1,
    )
    output_row_count = len(campaign_frame) + len(funnel_groups)

    workbook = load_workbook(template)
    report_ws = workbook[REPORT_SHEET]
    data_ws = workbook[DATA_SHEET]
    config_ws = workbook[CONFIG_SHEET]

    original_total_row = FIRST_DATA_ROW + TEMPLATE_CAPACITY
    if f"A{original_total_row}:E{original_total_row}" in {
        str(merged) for merged in report_ws.merged_cells.ranges
    }:
        report_ws.unmerge_cells(f"A{original_total_row}:E{original_total_row}")

    if output_row_count < TEMPLATE_CAPACITY:
        report_ws.delete_rows(
            FIRST_DATA_ROW + output_row_count,
            TEMPLATE_CAPACITY - output_row_count,
        )
    elif output_row_count > TEMPLATE_CAPACITY:
        report_ws.insert_rows(
            original_total_row,
            output_row_count - TEMPLATE_CAPACITY,
        )

    total_row = FIRST_DATA_ROW + output_row_count
    for row_number in range(FIRST_DATA_ROW, total_row):
        _copy_row_style(report_ws, FIRST_DATA_ROW, row_number)

    current_row = FIRST_DATA_ROW
    campaign_output_rows: dict[int, int] = {}
    for group_slug, funnel_value, group in funnel_groups:
        for _, row in group.iterrows():
            spending_daily = _as_number(row.get("stima_spending_giornaliera"))
            if spending_daily is None:
                spending_daily = _safe_divide(
                    row.get("investimento_media"),
                    dem_days_in_month if group_slug == "dem" else days_in_month,
                )
            values = [
                _as_text(row.get("funnel")),
                _as_text(row.get("platform")),
                _as_text(row.get("channel")),
                _campaign_id(row),
                _as_text(row.get("campaign_name")),
                _as_number(row.get("investimento_media")),
                _as_number(row.get("percentuale_investimento")),
                _as_number(row.get("cpp_medio")),
                _as_number(row.get("stima_pratiche")),
                _as_number(row.get("cpl_target")),
                _as_number(row.get("stima_lead")),
                _as_number(row.get("stima_lead_giornaliere")),
                _as_number(row.get("stima_lead_progressiva")),
                _as_number(row.get("lead_effettive")),
                _as_number(row.get("delta_lead")),
                spending_daily,
                _as_number(row.get("stima_spending_progressiva")),
                _as_number(row.get("speso_effettivo")),
                _as_number(row.get("delta_speso")),
                _as_number(row.get("delta_delivery_pct")),
                _as_number(row.get("cpl_effettivo")),
                _as_number(row.get("delta_cpl")),
                _as_action_text(row.get("action")),
            ]
            _write_report_values(report_ws, current_row, values)
            excel_row = _as_number(row.get("excel_row"))
            if excel_row is not None:
                campaign_output_rows[int(excel_row)] = current_row
            report_ws.cell(current_row, 5).alignment = Alignment(
                horizontal=report_ws.cell(current_row, 5).alignment.horizontal,
                vertical="top",
                wrap_text=True,
            )
            wrapped_lines = max(
                _wrapped_line_count(values[4]),
                _wrapped_line_count(values[-1]),
            )
            report_ws.row_dimensions[current_row].height = min(
                90,
                max(18, 6 + 15 * wrapped_lines),
            )
            current_row += 1

        subtotal_label = _as_text(funnel_value) or "Senza funnel"
        subtotal = official_frame[
            (official_frame.get("row_type", "") == "subtotal")
            & (official_frame["funnel"].astype(str) == f"TOT {subtotal_label}")
        ]
        subtotal_values = [f"TOT {subtotal_label}", None, None, None, None]
        if len(subtotal.index) == 1:
            subtotal_daily = campaign_daily_values.loc[group.index].dropna().sum()
            subtotal_values.extend(
                _official_summary_values(
                    subtotal.iloc[0], days_in_month, float(subtotal_daily)
                )
            )
        elif official_frame.empty:
            subtotal_values.extend(_legacy_summary_values(frame, group_slug))
        else:
            raise ValueError(f"Riga ufficiale TOT {subtotal_label} mancante o duplicata.")
        _write_report_values(report_ws, current_row, subtotal_values)
        report_ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=5)
        _style_subtotal(report_ws, current_row)
        current_row += 1

    _merge_combined_campaign_metrics(report_ws, campaign_output_rows)

    _copy_row_style(report_ws, total_row, total_row)
    official_total = official_frame[official_frame.get("row_type", "") == "total"]
    total_values = ["TOTALE GENERALE", None, None, None, None]
    if len(official_total.index) == 1:
        total_daily = campaign_daily_values.dropna().sum()
        total_values.extend(
            _official_summary_values(
                official_total.iloc[0], days_in_month, float(total_daily)
            )
        )
    elif official_frame.empty:
        total_values.extend(_legacy_summary_values(frame))
    else:
        raise ValueError("Riga ufficiale TOTALE GENERALE mancante o duplicata.")
    _write_report_values(report_ws, total_row, total_values)
    report_ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=5)
    _style_total(report_ws, total_row)

    period_text = (
        f"Periodo: {metadata['start_date']:%d/%m/%Y} - {metadata['end_date']:%d/%m/%Y}"
        f"  |  Ultimo aggiornamento: {metadata['updated_at']:%d/%m/%Y %H:%M}"
    )
    if metadata["source_status"] == "mock":
        period_text += "  |  DATI DEMO — NON INVIARE AL CLIENTE"
        report_ws["A2"].font = Font(
            name=report_ws["A2"].font.name,
            size=report_ws["A2"].font.sz,
            bold=True,
            color=RED_TEXT,
        )
    report_ws["A2"] = period_text

    _add_conditional_formatting(report_ws, total_row)
    report_ws.freeze_panes = "F5"
    report_ws.sheet_view.showGridLines = False
    report_ws.sheet_properties.pageSetUpPr.fitToPage = True
    report_ws.page_setup.orientation = "landscape"
    report_ws.page_setup.fitToWidth = 1
    report_ws.page_setup.fitToHeight = 0
    report_ws.page_margins.left = 0.25
    report_ws.page_margins.right = 0.25
    report_ws.page_margins.top = 0.5
    report_ws.page_margins.bottom = 0.5
    report_ws.print_title_rows = "4:4"
    report_ws.print_area = f"A1:W{total_row}"
    report_ws.print_options.horizontalCentered = True

    _write_data_sheet(data_ws, frame, days_in_month, dem_days_in_month)
    config_values = {
        "B2": metadata["client_name"],
        "B3": metadata["start_date"].date(),
        "B4": metadata["end_date"].date(),
        "B5": metadata["days_in_month"],
        "B6": metadata["elapsed_days"],
        "B7": metadata["updated_at"],
        "B8": metadata["source_status"],
    }
    for coordinate, value in config_values.items():
        config_ws[coordinate] = value
    config_ws["B3"].number_format = DATE_FORMAT
    config_ws["B4"].number_format = DATE_FORMAT
    config_ws["B7"].number_format = DATETIME_FORMAT
    config_ws["E2"] = "Lead effettive - Stima progressiva; positivo = sopra target"

    report_ws.sheet_state = "visible"
    data_ws.sheet_state = "hidden"
    config_ws.sheet_state = "hidden"
    workbook.active = workbook.sheetnames.index(REPORT_SHEET)

    destination_dir.mkdir(parents=True, exist_ok=True)
    dated_name = f"Report_Dynamica_{metadata['end_date']:%Y-%m-%d}.xlsx"
    dated_path = destination_dir / dated_name
    latest_path = destination_dir / LATEST_FILENAME
    temporary_path: Path | None = None
    latest_temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=".report_dynamica_",
            suffix=".xlsx",
            dir=destination_dir,
            delete=False,
        ) as temporary_file:
            temporary_path = Path(temporary_file.name)
        workbook.save(temporary_path)
        workbook.close()
        _validate_workbook(
            temporary_path, len(campaign_frame), metadata["end_date"].date()
        )
        os.replace(temporary_path, dated_path)
        temporary_path = None

        with tempfile.NamedTemporaryFile(
            prefix=".report_dynamica_latest_",
            suffix=".xlsx",
            dir=destination_dir,
            delete=False,
        ) as latest_temporary_file:
            latest_temporary_path = Path(latest_temporary_file.name)
        shutil.copyfile(dated_path, latest_temporary_path)
        os.replace(latest_temporary_path, latest_path)
        latest_temporary_path = None
    finally:
        workbook.close()
        for leftover in (temporary_path, latest_temporary_path):
            if leftover is not None:
                leftover.unlink(missing_ok=True)

    return str(dated_path)


def load_report_dataframe(
    report_path: str | os.PathLike[str] = "data/report_data.csv",
    metadata_path: str | os.PathLike[str] = "data/last_update.json",
) -> pd.DataFrame:
    """Load the final CSV and attach safe report metadata for the exporter."""
    report_file = _resolve_path(report_path, ROOT / "data/report_data.csv")
    metadata_file = _resolve_path(metadata_path, ROOT / "data/last_update.json")
    frame = pd.read_csv(
        report_file,
        dtype={
            "google_campaign_id": "string",
            "meta_campaign_id": "string",
            "dynamics_campaign_key": "string",
        },
    )
    if metadata_file.exists():
        frame.attrs["metadata"] = json.loads(metadata_file.read_text(encoding="utf-8"))
    return frame


def update_excel_export(
    report_path: Path = Path("data/report_data.csv"),
    output_path: Path = Path("exports/report_dynamica_updated.xlsx"),
) -> Path:
    """Backward-compatible wrapper around the client Excel generator."""
    frame = load_report_dataframe(report_path)
    generated = Path(generate_client_excel(frame, output_dir=str(output_path.parent)))
    latest = generated.parent / LATEST_FILENAME
    if output_path.name != LATEST_FILENAME:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(latest, output_path)
        return output_path
    return latest


def main() -> int:
    try:
        frame = load_report_dataframe()
        output_path = generate_client_excel(frame)
    except Exception as exc:
        print(f"Errore export Excel: {exc}")
        return 1
    print(f"Report Excel generato: {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
