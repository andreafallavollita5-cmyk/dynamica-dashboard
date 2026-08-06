"""Tests for the client Excel snapshot generated from the final DataFrame."""

from __future__ import annotations

import hashlib
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from src.update_excel import generate_client_excel


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "templates/Template_export_cliente_Dynamica.xlsx"


def report_row(**overrides) -> dict:
    row = {
        "report_date": "2026-07-13",
        "start_date": "2026-07-01",
        "end_date": "2026-07-12",
        "excel_row": 2,
        "funnel": "Funnel A",
        "platform": "Google",
        "channel": "Search",
        "campaign_name": "Campagna Google",
        "investimento_media": 1000.0,
        "percentuale_investimento": 0.5,
        "cpp_medio": 200.0,
        "stima_pratiche": 5.0,
        "cpl_target": 20.0,
        "stima_lead": 50.0,
        "stima_lead_giornaliere": 50 / 31,
        "stima_lead_progressiva": 600 / 31,
        "lead_effettive": 22.0,
        "delta_lead": 22 - (600 / 31),
        "stima_spending_progressiva": 12000 / 31,
        "speso_effettivo": 410.0,
        "delta_speso": 410 - (12000 / 31),
        "delta_delivery_pct": (410 - (12000 / 31)) / (12000 / 31),
        "cpl_effettivo": 410 / 22,
        "delta_cpl": (410 / 22) - 20,
        "action": "Ottimizzare creatività",
        "google_campaign_id": "12345678901234567890",
        "meta_campaign_id": "",
        "dynamics_campaign_key": "",
        "source_status": "google:ok",
    }
    row.update(overrides)
    return row


def make_frame(rows: list[dict], *, updated_at: str = "2026-07-13T07:00:00+02:00") -> pd.DataFrame:
    frame = pd.DataFrame(rows)
    frame.attrs["metadata"] = {
        "client": "Dynamica Retail",
        "start_date": "2026-07-01",
        "end_date": "2026-07-12",
        "updated_at": updated_at,
        "status": "ok",
    }
    return frame


def find_total_row(ws) -> int:
    return next(
        row for row in range(5, ws.max_row + 1)
        if ws.cell(row, 1).value == "TOTALE GENERALE"
    )


class ClientExcelExportTests(unittest.TestCase):
    def generate(self, frame: pd.DataFrame, output_dir: str) -> Path:
        return Path(
            generate_client_excel(
                frame,
                template_path=str(TEMPLATE),
                output_dir=output_dir,
            )
        )

    def test_mixed_platforms_missing_leads_zero_division_ids_actions_and_hidden_sheets(self):
        long_action = "Azione lunga " * 30
        frame = make_frame(
            [
                report_row(action=long_action),
                report_row(
                    platform="Meta",
                    channel="Social",
                    campaign_name="Campagna Meta",
                    google_campaign_id="",
                    meta_campaign_id="98765432109876543210",
                    lead_effettive=None,
                    cpl_effettivo=None,
                    delta_cpl=None,
                    source_status="meta:no_delivery",
                ),
                report_row(
                    platform="DEM",
                    channel="Email",
                    campaign_name="Campagna DEM",
                    google_campaign_id="",
                    dynamics_campaign_key="DEM-001",
                    lead_effettive=0,
                    cpl_effettivo=None,
                    delta_cpl=None,
                    action="",
                    source_status="manual:sheet_spend",
                ),
            ]
        )
        template_hash = hashlib.sha256(TEMPLATE.read_bytes()).hexdigest()
        with tempfile.TemporaryDirectory() as output_dir:
            path = self.generate(frame, output_dir)
            self.assertTrue(path.exists())
            self.assertEqual(path.read_bytes()[:2], b"PK")
            workbook = load_workbook(path, data_only=False)
            self.assertEqual(workbook.sheetnames, ["Report Cliente", "Dati", "Config"])
            self.assertEqual(workbook["Report Cliente"].sheet_state, "visible")
            self.assertEqual(workbook["Dati"].sheet_state, "hidden")
            self.assertEqual(workbook["Config"].sheet_state, "hidden")
            self.assertEqual(workbook["Config"]["B8"].value, "success")
            report = workbook["Report Cliente"]
            total_row = find_total_row(report)
            campaign_rows = [
                row for row in range(5, total_row)
                if report.cell(row, 5).value
            ]
            self.assertEqual(len(campaign_rows), 3)
            self.assertEqual(
                sum(1 for row in range(5, total_row) if str(report.cell(row, 1).value).startswith("TOT ")),
                2,
            )
            google_row = next(row for row in campaign_rows if report.cell(row, 5).value == "Campagna Google")
            self.assertEqual(report.cell(google_row, 4).value, "12345678901234567890")
            self.assertEqual(report.cell(google_row, 4).data_type, "s")
            self.assertEqual(report.cell(google_row, 23).value, long_action)
            meta_row = next(row for row in campaign_rows if report.cell(row, 5).value == "Campagna Meta")
            self.assertIsNone(report.cell(meta_row, 21).value)
            dem_row = next(row for row in campaign_rows if report.cell(row, 5).value == "Campagna DEM")
            self.assertAlmostEqual(report.cell(dem_row, 16).value, 1000 / 23)
            self.assertIsNone(report.cell(dem_row, 21).value)
            self.assertTrue(all(report.cell(row, 1).value or report.cell(row, 5).value for row in range(5, total_row)))
            data = workbook["Dati"]
            id_column = next(cell.column for cell in data[1] if cell.value == "google_campaign_id")
            campaign_column = next(cell.column for cell in data[1] if cell.value == "campaign_name")
            google_data_row = next(
                row for row in range(2, data.max_row + 1)
                if data.cell(row, campaign_column).value == "Campagna Google"
            )
            self.assertEqual(data.cell(google_data_row, id_column).value, "12345678901234567890")
            self.assertEqual(data.cell(google_data_row, id_column).number_format, "@")
            workbook.close()
        self.assertEqual(hashlib.sha256(TEMPLATE.read_bytes()).hexdigest(), template_hash)

    def test_multiple_funnels_create_subtotals_and_preserve_delta_convention(self):
        frame = make_frame(
            [
                report_row(
                    funnel="Area Clienti",
                    delta_lead=3.5,
                    kpi_delta_lead_totale=1.5,
                    subtotal_area_clienti_delta_lead=3.5,
                    subtotal_lead_veloce_delta_lead=-2.0,
                ),
                report_row(
                    funnel="Lead Veloce",
                    campaign_name="Seconda",
                    delta_lead=-2.0,
                ),
            ]
        )
        with tempfile.TemporaryDirectory() as output_dir:
            path = self.generate(frame, output_dir)
            workbook = load_workbook(path, data_only=True)
            report = workbook["Report Cliente"]
            total_row = find_total_row(report)
            subtotal_rows = [
                row for row in range(5, total_row)
                if str(report.cell(row, 1).value).startswith("TOT ")
            ]
            self.assertEqual([report.cell(row, 1).value for row in subtotal_rows], ["TOT Area Clienti", "TOT Lead Veloce"])
            self.assertAlmostEqual(report.cell(total_row, 15).value, 1.5)
            self.assertEqual(len(report.merged_cells.ranges), 5)
            workbook.close()

    def test_summary_rows_use_official_dataframe_fields_not_campaign_sums(self):
        frame = make_frame(
            [
                report_row(
                    funnel="Area Clienti",
                    kpi_investimento_media_totale=63000,
                    kpi_cpl_effettivo_totale=20.84,
                    subtotal_area_clienti_investimento_media=21000,
                    subtotal_area_clienti_cpl_effettivo=41.73,
                ),
                report_row(funnel="Area Clienti", campaign_name="Seconda"),
            ]
        )
        with tempfile.TemporaryDirectory() as output_dir:
            path = self.generate(frame, output_dir)
            workbook = load_workbook(path, data_only=True)
            report = workbook["Report Cliente"]
            total_row = find_total_row(report)
            subtotal_row = next(
                row for row in range(5, total_row)
                if report.cell(row, 1).value == "TOT Area Clienti"
            )
            self.assertEqual(report.cell(subtotal_row, 6).value, 21000)
            self.assertEqual(report.cell(subtotal_row, 21).value, 41.73)
            self.assertEqual(report.cell(total_row, 6).value, 63000)
            self.assertEqual(report.cell(total_row, 21).value, 20.84)
            workbook.close()

    def test_client_subtotals_are_area_clienti_lead_veloce_and_dem(self):
        frame = make_frame(
            [
                report_row(funnel="Area Clienti", campaign_name="Area"),
                report_row(funnel="Lead Veloce", campaign_name="Lead"),
                report_row(funnel="Pulsanti", campaign_name="Pulsanti"),
                report_row(funnel="Whatsapp", campaign_name="Whatsapp"),
                report_row(
                    funnel="Lead Veloce",
                    platform="Dem",
                    channel="DEM",
                    campaign_name="DEM",
                    google_campaign_id="",
                    source_status="manual:sheet_spend",
                ),
            ]
        )
        with tempfile.TemporaryDirectory() as output_dir:
            path = self.generate(frame, output_dir)
            workbook = load_workbook(path, data_only=True)
            report = workbook["Report Cliente"]
            total_row = find_total_row(report)
            subtotal_rows = [
                row for row in range(5, total_row)
                if str(report.cell(row, 1).value).startswith("TOT ")
            ]
            self.assertEqual(
                [report.cell(row, 1).value for row in subtotal_rows],
                ["TOT Area Clienti", "TOT Lead Veloce", "TOT DEM"],
            )
            lead_subtotal = subtotal_rows[1]
            dem_subtotal = subtotal_rows[2]
            self.assertLess(
                next(row for row in range(5, total_row) if report.cell(row, 5).value == "Pulsanti"),
                lead_subtotal,
            )
            self.assertLess(
                next(row for row in range(5, total_row) if report.cell(row, 5).value == "Whatsapp"),
                lead_subtotal,
            )
            self.assertGreater(
                next(row for row in range(5, total_row) if report.cell(row, 5).value == "DEM"),
                lead_subtotal,
            )
            self.assertLess(dem_subtotal, total_row)
            workbook.close()

    def test_more_than_thirty_campaigns_expand_before_total(self):
        rows = [
            report_row(campaign_name=f"Campagna {index:02d}", excel_row=index + 2)
            for index in range(31)
        ]
        with tempfile.TemporaryDirectory() as output_dir:
            path = self.generate(make_frame(rows), output_dir)
            workbook = load_workbook(path, data_only=False)
            report = workbook["Report Cliente"]
            total_row = find_total_row(report)
            self.assertEqual(total_row, 37)
            self.assertEqual(
                sum(1 for row in range(5, total_row) if report.cell(row, 5).value),
                31,
            )
            self.assertEqual(report.print_area, "'Report Cliente'!$A$1:$W$37")
            workbook.close()

    def test_single_funnel_under_capacity_has_no_stale_blank_rows(self):
        long_campaign = "Campagna molto lunga con nome descrittivo che deve occupare correttamente più righe nel report"
        with tempfile.TemporaryDirectory() as output_dir:
            path = self.generate(make_frame([report_row(campaign_name=long_campaign)]), output_dir)
            workbook = load_workbook(path, data_only=False)
            report = workbook["Report Cliente"]
            total_row = find_total_row(report)
            self.assertEqual(total_row, 7)
            self.assertEqual(report["E5"].value, long_campaign)
            self.assertGreaterEqual(report.row_dimensions[5].height, 45)
            self.assertEqual(report["A6"].value, "TOT Lead Veloce")
            workbook.close()

    def test_mock_status_adds_visible_warning(self):
        frame = make_frame([report_row(source_status="mock:demo")])
        with tempfile.TemporaryDirectory() as output_dir:
            path = self.generate(frame, output_dir)
            workbook = load_workbook(path, data_only=False)
            self.assertEqual(workbook["Config"]["B8"].value, "mock")
            self.assertIn("DATI DEMO — NON INVIARE AL CLIENTE", workbook["Report Cliente"]["A2"].value)
            workbook.close()

    def test_partial_status_and_repeated_generation_leave_no_previous_rows(self):
        first = make_frame(
            [report_row(campaign_name="Vecchia 1"), report_row(campaign_name="Vecchia 2")]
        )
        second = make_frame(
            [report_row(campaign_name="Nuova", source_status="google:error")]
        )
        with tempfile.TemporaryDirectory() as output_dir:
            self.generate(first, output_dir)
            self.generate(second, output_dir)
            latest = Path(output_dir) / "report_dynamica_updated.xlsx"
            workbook = load_workbook(latest, data_only=False)
            report = workbook["Report Cliente"]
            total_row = find_total_row(report)
            names = [report.cell(row, 5).value for row in range(5, total_row) if report.cell(row, 5).value]
            self.assertEqual(names, ["Nuova"])
            self.assertEqual(workbook["Config"]["B8"].value, "partial")
            workbook.close()

    def test_streamlit_download_is_connected_to_real_xlsx(self):
        source = (ROOT / "app.py").read_text(encoding="utf-8")
        self.assertIn("build_excel_download(", source)
        self.assertIn("excel_bytes,", source)
        self.assertIn("Report_Dynamica_{selected_start:%Y-%m-%d}_{selected_end:%Y-%m-%d}.xlsx", source)
        self.assertIn(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            source,
        )
        self.assertNotIn('file_name="report_data.csv"', source)

    def test_dashboard_table_has_excel_style_section_borders(self):
        source = (ROOT / "app.py").read_text(encoding="utf-8")
        self.assertIn("--table-section-border: #3a4454", source)
        self.assertIn("--table-section-border: #ffffff", source)
        for css_class in ("col-stima-spending", "col-cpl-target"):
            self.assertIn(f".campaign-table th.{css_class}", source)
            self.assertIn(f".campaign-table td.{css_class}", source)
        self.assertIn(".campaign-table th.col-campagna", source)
        self.assertIn(".campaign-table td.col-campagna", source)
        self.assertIn(
            "border-right:2px solid var(--table-section-border) !important",
            source,
        )
        self.assertIn(
            "border-left:2px solid var(--table-section-border) !important",
            source,
        )

    def test_adjacent_campaign_cells_are_not_merged(self):
        parent = report_row(
            excel_row=9,
            campaign_name="DYN_VELOCE Cessione del Quinto [Esatta]",
            lead_effettive=61,
            stima_lead_progressiva=128,
            delta_lead=25,
            cpl_effettivo=5490 / 153,
            delta_cpl=5490 / 153 - 38.5,
            cpl_target=38.5,
            speso_effettivo=5490,
        )
        child = report_row(
            excel_row=10,
            campaign_name=(
                "DYN_VELOCE Cessione del Quinto [Esatta] QUINTO DIGITALE"
            ),
            investimento_media=None,
            stima_lead=None,
            stima_lead_giornaliere=None,
            stima_lead_progressiva=None,
            lead_effettive=92,
            delta_lead=None,
            speso_effettivo=None,
            cpl_target=None,
            cpl_effettivo=None,
            delta_cpl=None,
        )
        with tempfile.TemporaryDirectory() as directory:
            latest = self.generate(make_frame([parent, child]), directory)
            workbook = load_workbook(latest, data_only=True)
            report = workbook["Report Cliente"]
            by_name = {
                report.cell(row, 5).value: row
                for row in range(5, report.max_row + 1)
                if report.cell(row, 5).value
            }
            first = by_name[parent["campaign_name"]]
            second = by_name[child["campaign_name"]]
            self.assertEqual(second, first + 1)
            merged = {str(item) for item in report.merged_cells.ranges}
            for column in ("K", "L", "M", "O", "U", "V"):
                self.assertNotIn(f"{column}{first}:{column}{second}", merged)
            self.assertNotIn(f"N{first}:N{second}", merged)
            self.assertEqual(report.cell(first, 14).value, 61)
            self.assertEqual(report.cell(second, 14).value, 92)
            workbook.close()


if __name__ == "__main__":
    unittest.main()
