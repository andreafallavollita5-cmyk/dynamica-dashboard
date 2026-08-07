from __future__ import annotations

import tempfile
import unittest
from io import BytesIO
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from app import (
    add_unmapped_daily_campaigns,
    apply_daily_spend_filter,
    build_dashboard_table_frame,
    build_csv_download,
    build_excel_download,
    build_period_report_frame,
    normalize_period,
    prepare_excel_export_frame,
)
from src.build_spend_daily import extract_daily_spend, validate_date_range


def report_rows() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "excel_row": 2,
                "report_date": "2026-07-04",
                "start_date": "2026-07-01",
                "end_date": "2026-07-03",
                "funnel": "Lead Veloce",
                "platform": "Google",
                "channel": "Search",
                "campaign_name": "Campagna Google",
                "google_campaign_id": "101",
                "meta_campaign_id": "",
                "investimento_media": 3100.0,
                "speso_effettivo": 999.0,
                "stima_spending_progressiva": 1200.0,
                "lead_effettive": 10.0,
                "stima_lead_progressiva": 20.0,
                "cpl_effettivo": 99.9,
                "delta_cpl": 79.9,
                "source_status": "google:ok",
            },
            {
                "excel_row": 3,
                "report_date": "2026-07-04",
                "start_date": "2026-07-01",
                "end_date": "2026-07-03",
                "funnel": "Lead Veloce",
                "platform": "Meta",
                "channel": "Display",
                "campaign_name": "Campagna Meta",
                "google_campaign_id": "",
                "meta_campaign_id": "202",
                "investimento_media": 6200.0,
                "speso_effettivo": 888.0,
                "stima_spending_progressiva": 2400.0,
                "lead_effettive": 4.0,
                "stima_lead_progressiva": 8.0,
                "cpl_effettivo": 222.0,
                "delta_cpl": 202.0,
                "source_status": "meta:ok",
            },
        ]
    )


def daily_rows() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {"date": date(2026, 7, 1), "source": "google_ads", "campaign_id": "101", "campaign_name": "G", "spend": 10.0},
            {"date": date(2026, 7, 2), "source": "google_ads", "campaign_id": "101", "campaign_name": "G", "spend": 20.0},
            {"date": date(2026, 7, 2), "source": "meta_ads", "campaign_id": "202", "campaign_name": "M", "spend": 7.0},
            {"date": date(2026, 7, 3), "source": "meta_ads", "campaign_id": "202", "campaign_name": "M", "spend": 0.0},
        ]
    )


class SpendDailyTests(unittest.TestCase):
    def test_scheduler_updates_daily_history_before_report(self):
        source = Path("run_daily_update.bat").read_text(encoding="utf-8")
        daily_position = source.index("-m src.build_spend_daily")
        report_position = source.index("-m src.build_report_data")
        self.assertLess(daily_position, report_position)

    def test_positive_unmapped_api_campaign_is_added_without_leads(self):
        daily = pd.concat([
            daily_rows(),
            pd.DataFrame([{
                "date": date(2026, 7, 2), "source": "meta_ads",
                "campaign_id": "303", "campaign_name": "Meta solo API", "spend": 9.0,
            }]),
        ], ignore_index=True)
        result = add_unmapped_daily_campaigns(pd, report_rows(), daily)
        added = result[result["meta_campaign_id"] == "303"].iloc[0]
        self.assertEqual(added["campaign_name"], "Meta solo API")
        self.assertEqual(added["funnel"], "API Ads")
        self.assertTrue(pd.isna(added["lead_effettive"]))
        self.assertTrue(pd.isna(added["cpl_effettivo"]))

    def test_full_range_and_estimate_are_dynamic(self):
        result = apply_daily_spend_filter(
            pd, report_rows(), daily_rows(), date(2026, 7, 1), date(2026, 7, 3)
        )
        self.assertEqual(result["speso_effettivo"].tolist(), [30.0, 7.0])
        self.assertEqual(result["stima_spending_progressiva"].tolist(), [300.0, 600.0])

    def test_dem_estimate_uses_only_weekdays_in_selected_range(self):
        report = report_rows().iloc[[0]].copy()
        report.loc[:, "platform"] = "DEM"
        report.loc[:, "channel"] = "DEM"
        report.loc[:, "investimento_media"] = 2300.0
        report.loc[:, "google_campaign_id"] = ""
        result = apply_daily_spend_filter(
            pd, report, daily_rows(), date(2026, 7, 1), date(2026, 7, 21)
        )
        self.assertEqual(result.iloc[0]["stima_spending_giornaliera"], 100.0)
        self.assertEqual(result.iloc[0]["stima_spending_progressiva"], 1500.0)

    def test_dem_weekend_only_range_has_zero_estimated_spend(self):
        report = report_rows().iloc[[0]].copy()
        report.loc[:, "platform"] = "DEM"
        report.loc[:, "channel"] = "DEM"
        report.loc[:, "google_campaign_id"] = ""
        result = apply_daily_spend_filter(
            pd, report, daily_rows(), date(2026, 7, 4), date(2026, 7, 5)
        )
        self.assertEqual(result.iloc[0]["stima_spending_progressiva"], 0.0)

    def test_single_day(self):
        result = apply_daily_spend_filter(
            pd, report_rows(), daily_rows(), date(2026, 7, 2), date(2026, 7, 2)
        )
        self.assertEqual(result["speso_effettivo"].tolist(), [20.0, 7.0])

    def test_multi_day(self):
        result = apply_daily_spend_filter(
            pd, report_rows(), daily_rows(), date(2026, 7, 1), date(2026, 7, 2)
        )
        self.assertEqual(result["speso_effettivo"].sum(), 37.0)

    def test_no_spend_interval_and_zero_spend_campaign(self):
        result = apply_daily_spend_filter(
            pd, report_rows(), daily_rows(), date(2026, 7, 3), date(2026, 7, 3)
        )
        self.assertEqual(result["speso_effettivo"].tolist(), [0.0, 0.0])

    def test_campaign_id_not_found_returns_zero(self):
        report = report_rows()
        report.loc[0, "google_campaign_id"] = "not-found"
        result = apply_daily_spend_filter(
            pd, report, daily_rows(), date(2026, 7, 1), date(2026, 7, 3)
        )
        self.assertEqual(result.loc[0, "speso_effettivo"], 0.0)

    def test_start_after_end_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "data iniziale"):
            validate_date_range(date(2026, 7, 2), date(2026, 7, 1))
        with self.assertRaises(ValueError):
            normalize_period((date(2026, 7, 2), date(2026, 7, 1)), date(2026, 7, 1), date(2026, 7, 3))

    def test_dates_from_different_months_are_rejected(self):
        with self.assertRaisesRegex(ValueError, "stesso mese"):
            normalize_period(
                (date(2026, 6, 30), date(2026, 7, 1)),
                date(2026, 6, 1),
                date(2026, 7, 15),
            )

    def test_one_api_error_is_controlled_and_other_data_is_written(self):
        def google_ok(start, end):
            return [{"date": start, "platform": "Google Ads", "campaign_id": "101", "campaign_name": "G", "spend": 12.0}]

        def meta_error(start, end):
            raise RuntimeError("Errore Meta controllato")

        with tempfile.TemporaryDirectory() as directory:
            output, statuses = extract_daily_spend(
                date(2026, 7, 1), date(2026, 7, 1), Path(directory) / "spend.csv",
                google_fetcher=google_ok, meta_fetcher=meta_error,
            )
            saved = pd.read_csv(output)
        self.assertEqual(len(saved.index), 1)
        self.assertEqual(statuses["Google Ads"]["status"], "ok")
        self.assertEqual(statuses["Meta Ads"]["status"], "error")

    def test_history_is_updated_and_deduplicated_by_required_key(self):
        def google_rows(start, end):
            return [
                {"date": start, "campaign_id": "101", "campaign_name": "G", "spend": 15},
                {"date": start, "campaign_id": "101", "campaign_name": "G", "spend": 16},
            ]

        def meta_rows(start, end):
            return []

        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "report_daily.csv"
            pd.DataFrame([
                {"date": "2026-06-30", "source": "google_ads", "campaign_id": "101", "campaign_name": "G", "spend": 9},
                {"date": "2026-07-01", "source": "google_ads", "campaign_id": "101", "campaign_name": "G", "spend": 10},
            ]).to_csv(output, index=False, encoding="utf-8-sig")
            extract_daily_spend(
                date(2026, 7, 1), date(2026, 7, 1), output,
                google_fetcher=google_rows, meta_fetcher=meta_rows,
            )
            saved = pd.read_csv(output, dtype={"campaign_id": str})
        self.assertEqual(len(saved), 2)
        self.assertEqual(saved.columns.tolist(), ["date", "source", "campaign_id", "campaign_name", "spend"])
        self.assertEqual(saved.loc[saved.date == "2026-07-01", "spend"].iloc[0], 16)

    def test_inactive_meta_campaign_is_excluded(self):
        inactive = "DYN_VELOCE LeadGen DIP - CQD | CBO Scaling (F3)"
        with tempfile.TemporaryDirectory() as directory:
            output, _ = extract_daily_spend(
                date(2026, 7, 1), date(2026, 7, 1), Path(directory) / "report_daily.csv",
                google_fetcher=lambda start, end: [],
                meta_fetcher=lambda start, end: [
                    {"date": start, "campaign_id": "999", "campaign_name": inactive, "spend": 50},
                    {"date": start, "campaign_id": "202", "campaign_name": "Attiva", "spend": 20},
                ],
            )
            saved = pd.read_csv(output, dtype={"campaign_id": str})
        self.assertEqual(saved.campaign_id.tolist(), ["202"])

    def test_recruitment_campaign_is_excluded_by_id_and_never_added_to_dashboard(self):
        recruitment_id = "6936446163375"
        with tempfile.TemporaryDirectory() as directory:
            output, _ = extract_daily_spend(
                date(2026, 7, 1), date(2026, 7, 1), Path(directory) / "report_daily.csv",
                google_fetcher=lambda start, end: [],
                meta_fetcher=lambda start, end: [
                    {
                        "date": start,
                        "campaign_id": recruitment_id,
                        "campaign_name": "Nome eventualmente modificato",
                        "spend": 147.64,
                    },
                    {
                        "date": start,
                        "campaign_id": "202",
                        "campaign_name": "Attiva",
                        "spend": 20,
                    },
                ],
            )
            saved = pd.read_csv(output, dtype={"campaign_id": str})
        self.assertNotIn(recruitment_id, saved.campaign_id.tolist())

        raw_daily = pd.DataFrame([{
            "date": date(2026, 7, 1),
            "source": "meta_ads",
            "campaign_id": recruitment_id,
            "campaign_name": "Recruitment Assicuratori 2026",
            "spend": 147.64,
        }])
        dashboard = add_unmapped_daily_campaigns(pd, report_rows(), raw_daily)
        self.assertNotIn(recruitment_id, dashboard["meta_campaign_id"].astype(str).tolist())

    def test_leads_and_cpl_columns_are_unchanged(self):
        original = report_rows()
        result = apply_daily_spend_filter(
            pd, original, daily_rows(), date(2026, 7, 2), date(2026, 7, 2)
        )
        for column in ("lead_effettive", "stima_lead_progressiva", "cpl_effettivo", "delta_cpl"):
            pd.testing.assert_series_equal(result[column], original[column])

    def test_quinto_digitale_continuation_is_consolidated(self):
        report = report_rows()
        continuation = report.iloc[0].copy()
        continuation["excel_row"] = 4
        continuation["campaign_name"] = "Campagna QUINTO DIGITALE"
        continuation["google_campaign_id"] = "102"
        continuation["investimento_media"] = None
        continuation["source_status"] = "google:rolled_up_to_excel_row_2"
        report = pd.concat([report, continuation.to_frame().T], ignore_index=True)
        daily = pd.concat([
            daily_rows(),
            pd.DataFrame([{
                "date": date(2026, 7, 2), "source": "google_ads",
                "campaign_id": "102", "campaign_name": "QUINTO", "spend": 5.0,
            }]),
        ], ignore_index=True)
        result = apply_daily_spend_filter(
            pd, report, daily, date(2026, 7, 2), date(2026, 7, 2)
        )
        self.assertEqual(result.loc[0, "speso_effettivo"], 25.0)
        self.assertTrue(pd.isna(result.loc[2, "speso_effettivo"]))

    def test_selected_period_updates_official_spending_summaries(self):
        result = apply_daily_spend_filter(
            pd, report_rows(), daily_rows(), date(2026, 7, 1), date(2026, 7, 3)
        )
        self.assertEqual(result["kpi_stima_spending_progressiva_totale"].iloc[0], 900)
        self.assertEqual(result["kpi_speso_effettivo_totale"].iloc[0], 37)
        self.assertEqual(result["kpi_delta_speso_totale"].iloc[0], -863)
        self.assertAlmostEqual(
            result["kpi_delta_delivery_pct_totale"].iloc[0], -863 / 900
        )
        self.assertEqual(
            result["subtotal_lead_veloce_speso_effettivo"].iloc[0], 37
        )

    def test_export_frame_uses_selected_dates_and_keeps_all_campaigns(self):
        metadata = {
            "client": "Dynamica Retail",
            "start_date": "2026-07-01",
            "end_date": "2026-07-03",
            "updated_at": "2026-07-04T07:00:00+02:00",
            "status": "ok",
        }
        export = prepare_excel_export_frame(
            pd,
            report_rows(),
            daily_rows(),
            metadata,
            date(2026, 7, 2),
            date(2026, 7, 2),
        )
        self.assertEqual(len(export), 4)
        self.assertEqual(
            export["row_type"].tolist(),
            ["campaign", "campaign", "subtotal", "total"],
        )
        self.assertEqual(export.attrs["metadata"]["start_date"], "2026-07-02")
        self.assertEqual(export.attrs["metadata"]["end_date"], "2026-07-02")
        self.assertEqual(export["start_date"].unique().tolist(), ["2026-07-02"])
        self.assertEqual(export["end_date"].unique().tolist(), ["2026-07-02"])
        self.assertEqual(export.iloc[:2]["speso_effettivo"].tolist(), [20.0, 7.0])

    def test_daily_crm_leads_update_leads_cpl_and_partial_official_rows(self):
        report = report_rows()
        report["row_type"] = "campaign"
        report["stima_lead"] = [310, 62]
        report["cpl_target"] = [20, 20]
        daily = daily_rows().copy()
        daily["excel_row"] = ""
        daily["leads"] = pd.NA
        daily["lead_allocation_method"] = ""
        daily = pd.concat(
            [
                daily,
                pd.DataFrame(
                    [
                        {
                            "date": date(2026, 7, 2), "source": "crm",
                            "excel_row": "2", "campaign_id": "",
                            "campaign_name": "Campagna Google", "spend": pd.NA,
                            "leads": 3, "lead_allocation_method": "crm_mapping",
                        },
                        {
                            "date": date(2026, 7, 2), "source": "crm",
                            "excel_row": "3", "campaign_id": "",
                            "campaign_name": "Campagna Meta", "spend": pd.NA,
                            "leads": 2, "lead_allocation_method": "crm_mapping",
                        },
                    ]
                ),
            ],
            ignore_index=True,
        )
        period = build_period_report_frame(
            pd, report, daily, date(2026, 7, 2), date(2026, 7, 2)
        )
        campaigns = period[period["row_type"] == "campaign"]
        self.assertEqual(campaigns["lead_effettive"].tolist(), [3.0, 2.0])
        self.assertEqual(campaigns["cpl_effettivo"].tolist(), [20 / 3, 3.5])
        self.assertEqual(
            period["row_type"].tolist(),
            ["campaign", "campaign", "subtotal", "total"],
        )
        subtotal = period[period["row_type"] == "subtotal"].iloc[0]
        total = period[period["row_type"] == "total"].iloc[0]
        self.assertEqual(subtotal["lead_effettive"], 5)
        self.assertEqual(total["speso_effettivo"], 27)

    def test_area_clienti_leads_stay_only_in_selected_period_total(self):
        report = pd.DataFrame(
            [
                {
                    **report_rows().iloc[0].to_dict(),
                    "excel_row": row,
                    "funnel": "Area Clienti",
                    "campaign_name": f"Area {row}",
                    "google_campaign_id": str(100 + row),
                }
                for row in range(2, 7)
            ]
        )
        daily = pd.DataFrame(
            [
                {
                    "date": date(2026, 7, 2), "source": "crm_area_clienti",
                    "excel_row": "", "campaign_id": "",
                    "campaign_name": "AREA CLIENTI", "spend": pd.NA,
                    "leads": 7, "lead_allocation_method": "area_clienti_total",
                }
            ]
        )
        result = build_period_report_frame(
            pd, report, daily, date(2026, 7, 2), date(2026, 7, 2)
        )
        campaigns = result[result["row_type"] == "campaign"]
        subtotal = result[
            (result["row_type"] == "subtotal")
            & (result["campaign_name"] == "TOT Area Clienti")
        ].iloc[0]
        total = result[result["row_type"] == "total"].iloc[0]
        self.assertTrue(campaigns["lead_effettive"].isna().all())
        self.assertEqual(subtotal["lead_effettive"], 7)
        self.assertEqual(total["lead_effettive"], 7)

    def test_downloaded_excel_matches_selected_period(self):
        metadata = {
            "client": "Dynamica Retail",
            "start_date": "2026-07-01",
            "end_date": "2026-07-03",
            "updated_at": "2026-07-04T07:00:00+02:00",
            "status": "ok",
        }
        content, filename = build_excel_download(
            pd,
            report_rows(),
            daily_rows(),
            metadata,
            date(2026, 7, 2),
            date(2026, 7, 2),
        )
        self.assertEqual(filename, "Report_Dynamica_2026-07-02_2026-07-02.xlsx")
        workbook = load_workbook(BytesIO(content), data_only=True)
        try:
            self.assertEqual(workbook["Config"]["B3"].value.date(), date(2026, 7, 2))
            self.assertEqual(workbook["Config"]["B4"].value.date(), date(2026, 7, 2))
            data = workbook["Dati"]
            start_column = next(
                cell.column for cell in data[1] if cell.value == "start_date"
            )
            end_column = next(
                cell.column for cell in data[1] if cell.value == "end_date"
            )
            self.assertEqual(
                data.cell(2, start_column).value.date(), date(2026, 7, 2)
            )
            self.assertEqual(
                data.cell(2, end_column).value.date(), date(2026, 7, 2)
            )
            report = workbook["Report Cliente"]
            campaign_rows = {
                report.cell(row, 5).value: row
                for row in range(5, report.max_row + 1)
                if report.cell(row, 5).value
            }
            self.assertEqual(set(campaign_rows), {"Campagna Google", "Campagna Meta"})
            self.assertEqual(
                report.cell(campaign_rows["Campagna Google"], 18).value, 20
            )
            total_row = next(
                row for row in range(5, report.max_row + 1)
                if report.cell(row, 1).value == "TOTALE GENERALE"
            )
            self.assertEqual(report.cell(total_row, 17).value, 300)
            self.assertEqual(report.cell(total_row, 18).value, 27)
        finally:
            workbook.close()

    def test_downloaded_csv_matches_selected_period(self):
        filtered = apply_daily_spend_filter(
            pd, report_rows(), daily_rows(), date(2026, 7, 2), date(2026, 7, 2)
        )
        content, filename = build_csv_download(
            filtered, date(2026, 7, 2), date(2026, 7, 2)
        )
        exported = pd.read_csv(BytesIO(content))
        self.assertEqual(filename, "report_data_2026-07-02_2026-07-02.csv")
        self.assertEqual(exported["start_date"].unique().tolist(), ["2026-07-02"])
        self.assertEqual(exported["end_date"].unique().tolist(), ["2026-07-02"])
        self.assertEqual(exported["speso_effettivo"].tolist(), [20.0, 7.0])

    def test_non_default_period_requires_daily_data(self):
        metadata = {"start_date": "2026-07-01", "end_date": "2026-07-03"}
        with self.assertRaisesRegex(ValueError, "dati giornalieri"):
            prepare_excel_export_frame(
                pd,
                report_rows(),
                pd.DataFrame(),
                metadata,
                date(2026, 7, 2),
                date(2026, 7, 2),
            )

    def test_dashboard_table_subtotal_and_total_follow_project_filter(self):
        all_rows = report_rows()
        static_filtered = all_rows.iloc[[0]].copy()
        dynamic_filtered = apply_daily_spend_filter(
            pd,
            static_filtered,
            daily_rows(),
            date(2026, 7, 1),
            date(2026, 7, 3),
        )
        table = build_dashboard_table_frame(
            pd,
            static_filtered,
            dynamic_filtered,
            all_rows,
            full_scope=False,
        )
        self.assertEqual(
            table["_row_type"].tolist(), ["campaign", "subtotal", "total"]
        )
        self.assertEqual(
            table["funnel"].tolist(),
            ["Lead Veloce", "TOT Lead Veloce", "TOTALE GENERALE"],
        )
        subtotal = table.iloc[1]
        total = table.iloc[2]
        self.assertEqual(subtotal["stima_spending_progressiva"], 300)
        self.assertEqual(subtotal["speso_effettivo"], 30)
        self.assertEqual(subtotal["lead_effettive"], 10)
        self.assertEqual(subtotal["cpl_effettivo"], 3)
        self.assertEqual(total["speso_effettivo"], 30)

    def test_row_type_report_keeps_recalculated_subtotal_with_project_filter(self):
        report = report_rows()
        report["row_type"] = "campaign"
        period = build_period_report_frame(
            pd, report, daily_rows(), date(2026, 7, 2), date(2026, 7, 2)
        )
        selected_static = report.iloc[[0]].copy()
        selected_dynamic = period[period["row_type"] == "campaign"].iloc[[0]].copy()
        table = build_dashboard_table_frame(
            pd,
            selected_static,
            selected_dynamic,
            period,
            full_scope=False,
        )
        self.assertEqual(
            table["_row_type"].tolist(), ["campaign", "subtotal", "total"]
        )
        self.assertEqual(table.iloc[1]["speso_effettivo"], 20)
        self.assertEqual(table.iloc[2]["speso_effettivo"], 20)

    def test_dashboard_table_uses_three_excel_business_subtotals(self):
        rows = report_rows()
        rows.loc[0, "funnel"] = "Area Clienti"
        dem = rows.iloc[[1]].copy()
        dem.index = [2]
        dem.loc[2, "funnel"] = "Lead Veloce"
        dem.loc[2, "platform"] = "Dem"
        dem.loc[2, "channel"] = "DEM"
        dem.loc[2, "campaign_name"] = "Campagna DEM"
        all_rows = pd.concat([rows, dem])
        dynamic = apply_daily_spend_filter(
            pd, all_rows, daily_rows(), date(2026, 7, 1), date(2026, 7, 3)
        )
        table = build_dashboard_table_frame(
            pd, all_rows, dynamic, all_rows, full_scope=True
        )
        labels = table.loc[
            table["_row_type"].isin(["subtotal", "total"]), "funnel"
        ].tolist()
        self.assertEqual(
            labels,
            [
                "TOT Area Clienti",
                "TOT Lead Veloce",
                "TOT DEM",
                "TOTALE GENERALE",
            ],
        )

    def test_dashboard_table_full_scope_keeps_official_manual_totals(self):
        static = report_rows()
        static["subtotal_lead_veloce_lead_effettive"] = 599
        static["kpi_lead_effettive_totale"] = 1338
        dynamic = apply_daily_spend_filter(
            pd, static, daily_rows(), date(2026, 7, 1), date(2026, 7, 3)
        )
        table = build_dashboard_table_frame(
            pd, static, dynamic, static, full_scope=True
        )
        subtotal = table[table["_row_type"] == "subtotal"].iloc[0]
        total = table[table["_row_type"] == "total"].iloc[0]
        self.assertEqual(subtotal["lead_effettive"], 599)
        self.assertEqual(total["lead_effettive"], 1338)
        self.assertEqual(subtotal["speso_effettivo"], 37)
        self.assertEqual(total["speso_effettivo"], 37)

    def test_dashboard_uses_official_totals_when_only_unplanned_api_rows_are_omitted(self):
        static = report_rows()
        official = static.iloc[[0]].copy()
        official.loc[:, "row_type"] = "total"
        official.loc[:, "funnel"] = "TOTALE GENERALE"
        official.loc[:, "campaign_name"] = "TOTALE GENERALE"
        official.loc[:, "cpl_target"] = 20.04
        unplanned = static.iloc[[0]].copy()
        unplanned["excel_row"] = None
        unplanned.loc[:, "funnel"] = "API Ads"
        unplanned.loc[:, "campaign_name"] = "Campagna API non pianificata"
        all_frame = pd.concat([static, unplanned, official], ignore_index=True)
        table = build_dashboard_table_frame(
            pd,
            static,
            static,
            all_frame,
            full_scope=False,
        )
        total = table[table["_row_type"] == "total"].iloc[0]
        self.assertEqual(total["cpl_target"], 20.04)


if __name__ == "__main__":
    unittest.main()
